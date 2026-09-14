"""DOCX parser built on python-docx."""

from __future__ import annotations

from pathlib import Path

import docx
from docx.document import Document as DocxDocument

from docintel.models import DocFormat, Document, Heading, Link


def _extract_hyperlinks(paragraph) -> list[Link]:
    """python-docx doesn't expose hyperlinks directly; read them from XML."""
    links: list[Link] = []
    rels = paragraph.part.rels
    for hyperlink in paragraph._p.findall(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}hyperlink"
    ):
        rid = hyperlink.get(
            "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
        )
        text = "".join(node.text or "" for node in hyperlink.iter() if node.tag.endswith("}t"))
        if rid and rid in rels:
            target = rels[rid].target_ref
            links.append(
                Link(text=text, target=target, is_external=target.startswith(("http", "mailto:")))
            )
    return links


def parse_docx(path: Path) -> Document:
    doc: DocxDocument = docx.Document(str(path))

    paragraphs = [p.text for p in doc.paragraphs]
    text = "\n".join(paragraphs)

    headings: list[Heading] = []
    links: list[Link] = []
    for p in doc.paragraphs:
        style = (p.style.name or "") if p.style is not None else ""
        if style.startswith("Heading"):
            try:
                level = int(style.removeprefix("Heading").strip())
            except ValueError:
                level = 1
            headings.append(Heading(level=level, text=p.text.strip()))
        links.extend(_extract_hyperlinks(p))

    metadata: dict[str, str] = {}
    core = doc.core_properties
    if core.title:
        metadata["title"] = core.title
    if core.author:
        metadata["author"] = core.author
    if core.subject:
        metadata["subject"] = core.subject

    return Document(
        path=path,
        format=DocFormat.DOCX,
        text=text,
        headings=headings,
        links=links,
        metadata=metadata,
    )
